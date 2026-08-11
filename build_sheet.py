"""
Builds Applications.xlsx - the per-family/client intake template for the
passport autofill bot. One file like this lives in Drive per client (linked
from the matching Zoho CRM deal). Columns = siblings/applicants in that same
family (so shared family data - address, emergency contact, parents - can be
typed once in column B and drag-filled across the other applicant columns).
Rows = fields, grouped into colored section bands. Columns/branching logic
are taken directly from notes.md.

Re-run this script any time the schema needs to change - it always
regenerates the whole file from scratch.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT_PATH = r"C:\Users\office_americandocs\Documents\AUTOFILL\Applications.xlsx"

NUM_APPLICANT_COLUMNS = 6  # siblings/applicants per family file; copy a column's formatting to add more

YES_NO = ["Yes", "No"]
SEX = ["Male", "Female"]
HAIR = ["BLACK", "BLONDE", "BROWN", "RED", "GRAY", "BALD", "OTHER"]
EYES = ["AMBER", "BLACK", "BLUE", "BROWN", "GRAY", "GREEN", "HAZEL"]
COMMUNICATION = ["Mail", "Email", "Both"]
STATUS = ["Needs Review", "Ready", "Done", "Error"]
PASSPORT_SCENARIO = ["First-time (None)", "Have Book", "Book Damaged", "Book Lost", "Book Stolen"]
NAME_CHANGE_TYPE = ["Marriage", "Court Order"]
HEIGHT_FEET = [str(n) for n in range(0, 11)]
HEIGHT_INCHES = [str(n) for n in range(0, 12)]
US_STATES = [
    "AL - ALABAMA", "AK - ALASKA", "AZ - ARIZONA", "AR - ARKANSAS", "CA - CALIFORNIA",
    "CO - COLORADO", "CT - CONNECTICUT", "DE - DELAWARE", "DC - DISTRICT OF COLUMBIA",
    "FL - FLORIDA", "GA - GEORGIA", "HI - HAWAII", "ID - IDAHO", "IL - ILLINOIS",
    "IN - INDIANA", "IA - IOWA", "KS - KANSAS", "KY - KENTUCKY", "LA - LOUISIANA",
    "ME - MAINE", "MD - MARYLAND", "MA - MASSACHUSETTS", "MI - MICHIGAN", "MN - MINNESOTA",
    "MS - MISSISSIPPI", "MO - MISSOURI", "MT - MONTANA", "NE - NEBRASKA", "NV - NEVADA",
    "NH - NEW HAMPSHIRE", "NJ - NEW JERSEY", "NM - NEW MEXICO", "NY - NEW YORK",
    "NC - NORTH CAROLINA", "ND - NORTH DAKOTA", "OH - OHIO", "OK - OKLAHOMA", "OR - OREGON",
    "PA - PENNSYLVANIA", "RI - RHODE ISLAND", "SC - SOUTH CAROLINA", "SD - SOUTH DAKOTA",
    "TN - TENNESSEE", "TX - TEXAS", "UT - UTAH", "VT - VERMONT", "VA - VIRGINIA",
    "WA - WASHINGTON", "WV - WEST VIRGINIA", "WI - WISCONSIN", "WY - WYOMING",
    "ASM - AMERICAN SAMOA", "GUM - GUAM", "XMI - MIDWAY ISLANDS", "MNP - NORTH MARIANA ISL.",
    "PRI - PUERTO RICO", "VIR - U.S. VIRGIN ISLANDS",
]

# Requirement tags (3rd element of each field tuple):
#   "R"   = always required, regardless of scenario
#   "O"   = always relevant but optional (never truly N/A)
#   "DYN" = required/not-relevant depends on other answers in the SAME
#           applicant column - colored dynamically via conditional
#           formatting (see build_google_sheet_template.py), not here
#   None  = system-managed field, not something staff fill in directly
#
# Section ORDER (2026-08-09 redesign): high-impact determining questions
# first so everything under them is understood to depend on them (Most
# Recent Passport -> Lost/Stolen -> Data Correction -> Parents & Spouse,
# in that cause-and-effect order), THEN sections that don't depend on any
# of that (Address, Emergency Contact, Travel, Other Names) - order among
# those last four doesn't matter, none of them gates the others.
#
# (section title, header fill color, [(row label, validation key or None, requirement)])
SECTIONS = [
    ("STATUS", "D9D9D9", [
        ("Status", "STATUS", None),
        ("Notes", None, None),
        ("Documents Link", None, "O"),
    ]),
    ("KEY QUESTIONS", "B45F06", [
        ("Date of Birth", None, "R"),
        ("Passport Scenario", "PASSPORT_SCENARIO", "R"),
        ("Book Issue Date", None, "DYN"),  # only if scenario != First-time
    ]),
    ("ABOUT YOU", "D9E1F2", [
        ("First Name", None, "R"),
        ("Middle Name", None, "O"),
        ("Last Name", None, "R"),
        ("Suffix", None, "O"),
        ("City of Birth", None, "R"),
        ("Country of Birth", None, "R"),
        ("State of Birth (USA only)", "STATE", "DYN"),  # required only if Country of Birth = USA
        ("SSN", None, "O"),
        ("USCIS A-Number", None, "O"),
        ("Sex", "SEX", "R"),
        ("Height Feet", "HEIGHT_FEET", "R"),
        ("Height Inches", "HEIGHT_INCHES", "R"),
        ("Hair Color", "HAIR", "R"),
        ("Eye Color", "EYES", "R"),
        ("Occupation", None, "R"),
        ("Employer", None, "O"),
    ]),
    ("MOST RECENT PASSPORT", "E4DFEC", [
        ("Reported Lost or Stolen?", "YES_NO", "DYN"),  # only if scenario = Book Lost/Stolen
        ("Limited Validity Under 2 Years?", "YES_NO", "DYN"),  # only if scenario != First-time
        ("Paid For Card Before?", "YES_NO", "DYN"),  # only if Limited Validity = Yes
        ("Name On Book - First", None, "DYN"),
        ("Name On Book - Last", None, "DYN"),
        ("Book Number", None, "DYN"),
    ]),
    ("LOST OR STOLEN REPORT (if Reported Lost or Stolen? = No)", "D6E4F0", [
        ("Reporting Own Passport?", "YES_NO", "DYN"),  # only if scenario=Lost/Stolen AND Reported=No
        ("Filed Police Report?", "YES_NO", "DYN"),
        ("Lost/Stolen Explanation", None, "DYN"),
        ("Lost/Stolen Location", None, "DYN"),
        ("Lost/Stolen Date", None, "DYN"),
        ("Other Passports Lost/Stolen?", "YES_NO", "DYN"),
    ]),
    ("DATA CORRECTION / DS-11 TRIGGER", "FFF2CC", [
        ("Data Printed Correctly?", "YES_NO", "DYN"),  # only if scenario != First-time
        ("Incorrect Fields", None, "DYN"),  # only if Data Printed Correctly = No
        ("Name Changed?", "YES_NO", "DYN"),  # only if scenario != First-time
        ("Name Change Type", "NAME_CHANGE_TYPE", "DYN"),  # only if Name Changed = Yes
        ("Name Change Place", None, "DYN"),  # City/State - only if Name Changed = Yes
        ("Name Change Date", None, "DYN"),  # only if Name Changed = Yes
        ("Name Change Certified Docs?", "YES_NO", "DYN"),  # only if Name Changed = Yes - No means DS-11 required
    ]),
    ("PARENTS & SPOUSE (DS-11 ONLY)", "F8CBAD", [
        ("Parent 1 Unknown?", "YES_NO", "DYN"),  # only if DS-11 path likely (see notes.md)
        ("Parent 1 First & Middle Name", None, "DYN"),  # + Parent 1 Unknown != Yes
        ("Parent 1 Last Name", None, "DYN"),
        ("Parent 1 Date of Birth", None, "DYN"),
        ("Parent 1 Place of Birth", None, "DYN"),
        ("Parent 1 Sex", "SEX", "DYN"),
        ("Parent 1 US Citizen?", "YES_NO", "DYN"),
        ("Parent 2 Unknown?", "YES_NO", "DYN"),
        ("Parent 2 First & Middle Name", None, "DYN"),
        ("Parent 2 Last Name", None, "DYN"),
        ("Parent 2 Date of Birth", None, "DYN"),
        ("Parent 2 Place of Birth", None, "DYN"),
        ("Parent 2 Sex", "SEX", "DYN"),
        ("Parent 2 US Citizen?", "YES_NO", "DYN"),
        ("Ever Married?", "YES_NO", "DYN"),
        ("Spouse First & Middle Name", None, "DYN"),  # + Ever Married = Yes
        ("Spouse Last Name", None, "DYN"),
        ("Spouse Date of Birth", None, "DYN"),
        ("Spouse Place of Birth", None, "DYN"),
        ("Spouse US Citizen?", "YES_NO", "DYN"),
        ("Marriage Date", None, "DYN"),
        ("Ever Divorced or Widowed?", "YES_NO", "DYN"),
        ("Divorce Date", None, "DYN"),  # + Ever Divorced or Widowed = Yes
    ]),
    ("MAILING ADDRESS", "E2EFDA", [
        ("Mail Street", None, "R"),
        ("Mail City", None, "R"),
        ("Mail Country", None, "R"),
        ("Mail State (USA only)", "STATE", "DYN"),  # only if Mail Country = USA
        ("Mail Zip", None, "DYN"),  # only if Mail Country = USA
        ("In Care Of", None, "O"),
        ("Same As Permanent Address?", "YES_NO", "R"),
        ("Permanent Street", None, "DYN"),  # only if Same As Permanent = No
        ("Permanent Apartment", None, "DYN"),
        ("Permanent City", None, "DYN"),
        ("Permanent Country", None, "DYN"),
        ("Permanent State (USA only)", "STATE", "DYN"),  # + Permanent Country = USA
        ("Permanent Zip", None, "DYN"),  # + Permanent Country = USA
        ("Preferred Communication", "COMMUNICATION", "R"),
        ("Email", None, "R"),
        ("Confirm Email", None, "R"),
    ]),
    ("EMERGENCY CONTACT", "FCE4D6", [
        ("EC Name", None, "R"),
        ("EC Address", None, "R"),
        ("EC Apartment", None, "O"),
        ("EC City", None, "R"),
        ("EC Country", None, "R"),
        ("EC State (USA only)", "STATE", "DYN"),  # only if EC Country = USA
        ("EC Zip", None, "R"),  # required even for non-US addresses - see run_autofill.py us_zip_prefix()
        ("EC Phone", None, "R"),
        ("EC Email", None, "R"),  # confirmed live 2026-08-10: site marks it required, rejects blank
        ("EC Relationship", None, "R"),
    ]),
    ("TRAVEL PLANS", "D6E4F0", [
        ("Trip Date", None, "O"),
        ("Trip Return Date", None, "O"),
        ("Countries To Be Visited", None, "O"),
    ]),
    ("OTHER NAMES", "D9D2E9", [
        ("Other Name 1 - First", None, "O"),
        ("Other Name 1 - Last", None, "O"),
        ("Other Name 2 - First", None, "O"),
        ("Other Name 2 - Last", None, "O"),
        ("Other Name 3 - First", None, "O"),
        ("Other Name 3 - Last", None, "O"),
    ]),
]

VALIDATION_LISTS = {
    "YES_NO": YES_NO,
    "SEX": SEX,
    "HAIR": HAIR,
    "EYES": EYES,
    "COMMUNICATION": COMMUNICATION,
    "STATUS": STATUS,
    "PASSPORT_SCENARIO": PASSPORT_SCENARIO,
    "NAME_CHANGE_TYPE": NAME_CHANGE_TYPE,
    "HEIGHT_FEET": HEIGHT_FEET,
    "HEIGHT_INCHES": HEIGHT_INCHES,
    "STATE": US_STATES,
}


def build():
    wb = openpyxl.Workbook()

    # ---- hidden Lists sheet (data validation sources) ----
    lists_ws = wb.create_sheet("Lists")
    lists_ws.sheet_state = "hidden"
    list_ranges = {}
    for col_idx, (key, values) in enumerate(VALIDATION_LISTS.items(), start=1):
        col_letter = get_column_letter(col_idx)
        lists_ws.cell(row=1, column=col_idx, value=key)
        for r, v in enumerate(values, start=2):
            lists_ws.cell(row=r, column=col_idx, value=v)
        list_ranges[key] = f"Lists!${col_letter}$2:${col_letter}${len(values) + 1}"

    # ---- main Applicants sheet (rows = fields, columns = siblings) ----
    ws = wb.active
    ws.title = "Applicants"

    last_col_idx = 1 + NUM_APPLICANT_COLUMNS  # column A = labels, then B..
    last_col_letter = get_column_letter(last_col_idx)

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.column_dimensions["A"].width = 30
    for c in range(2, last_col_idx + 1):
        ws.column_dimensions[get_column_letter(c)].width = 20

    # applicant column headers (row 1): "Applicant 1", "Applicant 2", ...
    header_row = 1
    label_cell = ws.cell(row=header_row, column=1, value="APPLICANT ->")
    label_cell.font = Font(bold=True, size=10, italic=True, color="808080")
    label_cell.alignment = Alignment(horizontal="right", vertical="center")
    for i in range(NUM_APPLICANT_COLUMNS):
        c = ws.cell(row=header_row, column=2 + i, value=f"Applicant {i + 1}")
        c.font = Font(bold=True, size=11)
        c.fill = PatternFill("solid", fgColor="404040")
        c.font = Font(bold=True, size=11, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[header_row].height = 20

    row = header_row + 1
    for section_title, color_hex, fields in SECTIONS:
        # section title band, merged across label + all applicant columns
        band = ws.cell(row=row, column=1, value=section_title)
        band.font = Font(bold=True, size=11, color="FFFFFF")
        band.fill = PatternFill("solid", fgColor=color_hex)
        band.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col_idx)
        for cc in range(1, last_col_idx + 1):
            ws.cell(row=row, column=cc).fill = PatternFill("solid", fgColor=color_hex)
        ws.row_dimensions[row].height = 20
        row += 1

        for label, validation_key, _requirement in fields:
            label_c = ws.cell(row=row, column=1, value=label)
            label_c.font = Font(bold=True, size=10)
            label_c.fill = PatternFill("solid", fgColor=color_hex, tint=0.6) if False else PatternFill("solid", fgColor=_lighten(color_hex))
            label_c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
            label_c.border = border

            for cc in range(2, last_col_idx + 1):
                data_c = ws.cell(row=row, column=cc)
                data_c.border = border
                data_c.alignment = Alignment(horizontal="center", vertical="center")

            if validation_key:
                dv = DataValidation(
                    type="list",
                    formula1=list_ranges[validation_key],
                    allow_blank=True,
                    showDropDown=False,
                )
                dv.error = "Please choose a value from the dropdown list."
                dv.errorTitle = "Invalid entry"
                ws.add_data_validation(dv)
                dv.add(f"B{row}:{last_col_letter}{row}")

            row += 1

    ws.freeze_panes = "B2"  # keep field labels (col A) + applicant header row (1) visible
    ws.sheet_view.showGridLines = True

    wb.save(OUT_PATH)
    print(f"Saved {OUT_PATH}")
    print(f"Total rows: {row - 1}, applicant columns: {NUM_APPLICANT_COLUMNS}")


def _lighten(hex_color, factor=0.55):
    """Lighten a hex color toward white for the field-label cells (softer than the section band)."""
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    r = int(r + (255 - r) * factor)
    g = int(g + (255 - g) * factor)
    b = int(b + (255 - b) * factor)
    return f"{r:02X}{g:02X}{b:02X}"


if __name__ == "__main__":
    build()
